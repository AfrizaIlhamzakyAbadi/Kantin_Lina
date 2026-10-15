import docx

doc = docx.Document('Report/Report.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith(('Bab 5', 'Bab 6', 'Bab 7', 'Bab 8', 'Bab 9', 'Lampiran')):
        print(f"Paragraph {i}: [{p.style.name}] {txt}")
