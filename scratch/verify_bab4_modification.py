import docx

doc = docx.Document('Report/Report.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total inline shapes: {len(doc.inline_shapes)}")

print("\n=== CHAPTER HEADINGS CHECK ===")
chapters = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name == 'Heading 1' or txt.startswith(('Bab', 'Lampiran')):
        print(f"P[{i}] [{p.style.name}]: {txt}")
        chapters.append((i, txt))

print("\n=== BAB 4 CONTENT BREAKDOWN ===")
bab4_start = None
bab5_start = None
for i, txt in chapters:
    if txt.startswith("Bab 4"):
        bab4_start = i
    elif txt.startswith("Bab 5"):
        bab5_start = i

print(f"Bab 4 starts at paragraph {bab4_start}, Bab 5 starts at paragraph {bab5_start}")

bab4_drawings = 0
for idx in range(bab4_start, bab5_start):
    p = doc.paragraphs[idx]
    drawings = len(p._p.xpath('.//w:drawing'))
    bab4_drawings += drawings
    if p.style.name.startswith('Heading') or drawings > 0 or p.text.strip().startswith(('Tabel', 'Gambar')):
        print(f"P[{idx}] [{p.style.name}] (drawings={drawings}): {p.text.strip()[:80]}")

print(f"\nTotal drawings in Bab 4: {bab4_drawings}")

print("\n=== TABLES IN BAB 4 ===")
for t_idx, tbl in enumerate(doc.tables):
    # Check if table belongs to Bab 4
    first_cell = tbl.rows[0].cells[0].text.strip()
    if any(k in first_cell for k in ["Kode Use Case", "Notasi UML", "Komponen Diagram", "Nama Simbol"]):
        print(f"Table {t_idx} ({len(tbl.rows)} rows x {len(tbl.columns)} cols): Header: {[c.text.strip() for c in tbl.rows[0].cells]}")
        print(f"   Row 1: {[c.text.strip() for c in tbl.rows[1].cells]}")
