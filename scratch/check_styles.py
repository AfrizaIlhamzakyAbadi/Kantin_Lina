import docx

doc = docx.Document('Report/Report.docx')

for i in [49, 50, 51, 60, 69, 70, 71]:
    p = doc.paragraphs[i]
    font_name = p.runs[0].font.name if p.runs and p.runs[0].font.name else 'Default'
    font_size = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 'Default'
    bold = p.runs[0].bold if p.runs else False
    print(f"P[{i}] style='{p.style.name}' font='{font_name}' size={font_size} bold={bold} text='{p.text[:40]}'")

# Check table styling
tbl = doc.tables[11]
print(f"Table 11 style: {tbl.style.name}")
