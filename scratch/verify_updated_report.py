import docx

doc = docx.Document('Report/Report.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total inline shapes: {len(doc.inline_shapes)}")

print("\n=== HEADINGS AND DRAWINGS ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    drawings = p._p.xpath('.//w:drawing')
    if p.style.name.startswith('Heading') or txt.startswith(('Bab', 'Lampiran')):
        print(f"P[{i}] [{p.style.name}] {txt[:70]}")
    elif drawings:
        nxt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
        print(f"P[{i}] DRAWING({len(drawings)}) -> Caption: {nxt[:70]}")

print("\n=== VERIFYING BAB 5 TABLES ===")
for t_idx in range(11, len(doc.tables)):
    tbl = doc.tables[t_idx]
    print(f"Table {t_idx} ({len(tbl.rows)} rows x {len(tbl.columns)} cols): Header: {[c.text.strip() for c in tbl.rows[0].cells]}")
    if len(tbl.rows) > 1:
        print(f"   Row 1: {[c.text.strip() for c in tbl.rows[1].cells]}")
