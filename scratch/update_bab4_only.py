import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Paths
doc_path = r"c:\Users\ACER SWIFT\AppProject\Kantin Lina v0.0\Report\Report.docx"
backup_path = r"c:\Users\ACER SWIFT\AppProject\Kantin Lina v0.0\Report\Report_backup_before_bab4_update.docx"
diagrams_dir = r"c:\Users\ACER SWIFT\AppProject\Kantin Lina v0.0\Report\diagrams"

# 1. Create a safe backup before modifying
shutil.copy2(doc_path, backup_path)
print(f"Created safety backup: {backup_path}")

doc = docx.Document(doc_path)

# Locate boundaries: Bab 4 start and Bab 5 start in body._element
body_children = doc._body._element[:]
bab4_start_idx = None
bab5_start_idx = None

for idx, child in enumerate(body_children):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(child, doc)
        txt = p.text.strip()
        if txt.startswith("Bab 4") and bab4_start_idx is None:
            bab4_start_idx = idx
        elif txt.startswith("Bab 5") and bab5_start_idx is None:
            bab5_start_idx = idx

print(f"Bab 4 element index: {bab4_start_idx}")
print(f"Bab 5 element index: {bab5_start_idx}")

if bab4_start_idx is None or bab5_start_idx is None:
    raise ValueError("Could not find Bab 4 or Bab 5 boundaries in document.")

# Reference element before which new content will be inserted (Bab 5 start)
bab5_elem = body_children[bab5_start_idx]

# Remove old elements between Bab 4 and Bab 5
for idx in range(bab4_start_idx, bab5_start_idx):
    elem_to_remove = body_children[idx]
    doc._body._element.remove(elem_to_remove)

print(f"Removed {bab5_start_idx - bab4_start_idx} old elements of Bab 4.")

# Helper functions to create formatted elements and insert before bab5_elem
def create_p(text="", style='Normal', space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        if style == 'Heading 1':
            run.font.size = Pt(16)
            run.bold = True
        elif style == 'Heading 2':
            run.font.size = Pt(13)
            run.bold = True
        elif style == 'Heading 3':
            run.font.size = Pt(11.5)
            run.bold = True
        else:
            run.font.size = Pt(11)
    doc._body._element.remove(p._p)
    bab5_elem.addprevious(p._p)
    return p

def create_bullet(text, bold_prefix="", space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        r1.bold = True
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    doc._body._element.remove(p._p)
    bab5_elem.addprevious(p._p)
    return p

def create_caption(text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    r.font.italic = True
    doc._body._element.remove(p._p)
    bab5_elem.addprevious(p._p)
    return p

def create_image(image_filename, caption_text, width_inches=6.0):
    img_path = os.path.join(diagrams_dir, image_filename)
    if not os.path.exists(img_path):
        print(f"Warning: image file not found: {img_path}")
        return
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    doc._body._element.remove(p._p)
    bab5_elem.addprevious(p._p)
    create_caption(caption_text)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'CCCCCC'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def create_table(headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = tbl.rows[0].cells
    for c_idx, title in enumerate(headers):
        cell = hdr_cells[c_idx]
        cell.text = title
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        if p.runs:
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(10)
            p.runs[0].bold = True
        set_cell_shading(cell, "F2F2F2")
        set_cell_border(cell, top={'sz': 6, 'color': '888888'}, bottom={'sz': 6, 'color': '888888'},
                              left={'sz': 4, 'color': 'CCCCCC'}, right={'sz': 4, 'color': 'CCCCCC'})
    
    # Body rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = "FAFAFA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row_cells[c_idx]
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if p.runs:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(9.5)
            set_cell_shading(cell, bg_color)
            set_cell_border(cell, top={'sz': 4, 'color': 'E0E0E0'}, bottom={'sz': 4, 'color': 'E0E0E0'},
                                  left={'sz': 4, 'color': 'E0E0E0'}, right={'sz': 4, 'color': 'E0E0E0'})
    
    if col_widths:
        for r in tbl.rows:
            for idx, width in enumerate(col_widths):
                r.cells[idx].width = Inches(width)

    doc._body._element.remove(tbl._tbl)
    bab5_elem.addprevious(tbl._tbl)
    # spacer paragraph
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(6)
    doc._body._element.remove(p_space._p)
    bab5_elem.addprevious(p_space._p)

print("Writing newly structured Bab 4...")

# ==============================================================================
# BAB 4: ACTIVITY DIAGRAM
# ==============================================================================
create_p("Bab 4 Activity Diagram", style='Heading 1', space_before=12, space_after=6)

create_p(
    "Activity Diagram (Diagram Aktivitas) merupakan model visual berbasis UML (Unified Modeling Language) "
    "yang menggambarkan aliran kerja prosedural (workflow), perilaku bisnis, dan sekuens logis aktivitas "
    "dari sistem yang direkayasa balik (reverse engineered). Berdasarkan panduan Modul 4 Praktikum APSI, "
    "Activity Diagram memfokuskan pemodelan pada aktivitas operasional yang dieksekusi oleh sistem dan komponen "
    "pendukungnya dalam merealisasikan skenario Use Case yang telah didefinisikan pada Bab 3.",
    space_after=6
)

create_p(
    "Struktur diagram ini mengadopsi partisi vertikal (Swimlane) guna mempertegas batas tanggung jawab (separation "
    "of concerns) antar entitas. Dalam Sistem Informasi Pre-Order Kantin Lina, partisi melibatkan Aktor Pengguna "
    "(Pelanggan dan Admin Kantin), Edge Compute Serverless (Cloudflare Worker), Layanan Pihak Ketiga (Midtrans "
    "Payment Gateway), serta Basis Data Sinkronisasi Reaktif (Cloud Firestore). Tabel 4.1 merangkum pemetaan antara "
    "Use Case pada Bab 3 dengan empat Activity Diagram utama yang dimodelkan pada bab ini.",
    space_after=6
)

# Table 4.1: Mapping Use Case to Activity Diagram
t41_headers = ["No", "Kode Use Case", "Nama Use Case (Bab 3)", "Activity Diagram Terkait"]
t41_data = [
    ["1", "UC01, UC02, UC03, UC04", "Melihat Menu, Memilih Slot Jam, Pembayaran QRIS, Pemantauan Tiket", "4.1 Alur Pre-Order & Checkout QRIS Dinamis"],
    ["2", "UC03, UC04, UC08, UC10", "Pembayaran QRIS, Pemantauan Tiket, Antrean Prioritas, Webhook Notifikasi", "4.2 Alur Verifikasi Webhook Midtrans & Penomoran Antrean"],
    ["3", "UC04, UC08, UC09", "Pemantauan Status Pesanan, Monitor Antrean Prioritas, Ubah Status Selesai", "4.3 Alur Penyiapan Dapur & Penyerahan Makanan di Loket"],
    ["4", "UC05, UC06", "Login Admin, Pengelolaan Katalog Menu & Stok Harian (Food/Drink/Other)", "4.4 Alur Pengelolaan Katalog Menu & Stok Harian"],
]
create_table(t41_headers, t41_data, col_widths=[0.5, 1.3, 2.7, 2.0])

create_p(
    "Seluruh diagram disusun mengikuti standar notasi Modul 4 Praktikum APSI yang terdiri dari tujuh simbol baku, "
    "sebagaimana diikhtisarkan pada Tabel 4.2.",
    space_after=6
)

# Table 4.2: Activity Diagram Standard Notation (Module 4)
t42_headers = ["Nama Simbol", "Notasi UML", "Fungsi & Penerapan dalam Sistem Kantin Lina"]
t42_data = [
    ["Start Point (Initial Node)", "Lingkaran solid hitam", "Menandai titik awal workflow pemrosesan (contoh: pelanggan membuka web atau webhook diterima)."],
    ["Activity / Action Node", "Persegi panjang sudut tumpul (lozenge)", "Menggambarkan unit tugas operasional sistem atau interaksi fisik pengguna."],
    ["State Transition", "Garis solid berpanah terarah", "Menunjukkan arah sekuens eksekusi aktivitas dari satu state ke state selanjutnya."],
    ["Decision & Merge Node", "Simbol belah ketupat (diamond)", "Mengevaluasi kondisi percabangan (guard condition) seperti validasi stok, signature, dan TTL 1 menit."],
    ["Fork & Join", "Batang horizontal/vertikal padat", "Menunjukkan percabangan eksekusi paralel dan penggabungan sinkronisasi beberapa alur bersamaan."],
    ["End Point (Final Node)", "Lingkaran ganda / bull's eye", "Menggambarkan terminasi akhir dari keseluruhan alur workflow aktivitas."],
    ["Swimlane (Partition)", "Kolom pembagi vertikal", "Memisahkan tanggung jawab fungsional antara Pelanggan, Worker, Midtrans, Firestore, dan Admin."],
]
create_table(t42_headers, t42_data, col_widths=[1.5, 1.5, 3.5])

# ------------------------------------------------------------------------------
# 4.1 Activity Diagram Pemesanan & Checkout QRIS Dinamis
# ------------------------------------------------------------------------------
create_p("4.1 Activity Diagram Pemesanan & Checkout QRIS Dinamis", style='Heading 2', space_before=10, space_after=4)

create_p(
    "Diagram aktivitas pada Gambar 4.1 merepresentasikan alur utama pemesanan makanan secara mandiri oleh pelanggan "
    "(siswa dan guru). Sesuai filosofi Zero-Friction Guest Ordering yang diterapkan pada kode sumber active codebase, "
    "pelanggan dapat langsung bertransaksi tanpa perlu melakukan pendaftaran akun ataupun login kredensial.",
    space_after=6
)

create_p(
    "Alur ini memodelkan validasi ketat pada sisi klien dan server: pelanggan memilih menu dari katalog terkelola, "
    "memasukkan nama panggilan (pickup name), serta menentukan jadwal pengambilan berbasis batch 5 menit dengan "
    "lead time buffer minimal 2 menit. Saat tombol checkout ditekan, Cloudflare Edge Worker (/api/checkout) secara "
    "atomik mengeksekusi Firestore Transaction untuk memotong stok (Soft Allocation), mencatat pesanan 'Pending' "
    "dengan deadline pembayaran ketat 1 menit (60 detik, PAYMENT_WINDOW_MS = 60000, EXPIRY_GRACE_MS = 0), dan "
    "menginisialisasi Dynamic QRIS melalui Midtrans Core API. Perlu ditekankan bahwa nomor urut antrean harian "
    "(order_tag) sengaja TIDAK dibangkitkan pada tahap ini guna menjamin nomor antrean hanya diberikan kepada pesanan "
    "yang telah terverifikasi lunas.",
    space_after=6
)

# Table 4.3: Specification of Activity Diagram 4.1
t43_headers = ["Komponen Diagram", "Elemen Pemodelan", "Deskripsi Perilaku dalam Sistem"]
t43_data = [
    ["Partisi Swimlane", "Pelanggan (Web PWA)", "Menavigasi menu, mengatur kuota keranjang, mengisi nama, memilih slot jam & menit, memindai QRIS."],
    ["Partisi Swimlane", "Cloudflare Worker (/api/checkout)", "Memvalidasi payload, mengeksekusi transaksi Firestore atomik, memanggil Midtrans Core API."],
    ["Partisi Swimlane", "Midtrans Payment Gateway", "Memvalidasi parameter tagihan, menerbitkan Dynamic QRIS string dengan masa berlaku 1 menit."],
    ["Partisi Swimlane", "Layar Pembayaran Pelanggan", "Merender citra QRIS dinamis, memasang listener onSnapshot, menjalankan hitung mundur 60 detik."],
    ["Decision Point", "Stok Tersedia & Cukup?", "Jika stok = 0 atau kurang dari pesanan, transaksi dibatalkan dengan HTTP 409 Conflict. Jika cukup, stok dikurangi atomik."],
    ["Guard Condition", "Validasi Charge Midtrans", "Jika Midtrans menerima request (HTTP 200/201), QRIS diteruskan ke klien. Jika ditolak, transaksi dibatalkan."],
]
create_table(t43_headers, t43_data, col_widths=[1.5, 1.8, 3.2])

create_p("Alur eksekusi rinci terbagi ke dalam skenario normal dan skenario penanganan eksepsi sebagai berikut:", space_after=4)
create_bullet("Pelanggan mengakses situs web Kantin Lina pada peramban smartphone tanpa hambatan autentikasi.", bold_prefix="1. Akses Katalog: ")
create_bullet("Pelanggan memilih item makanan/minuman dan menentukan kuota pemesanan (maksimal 10 item per jenis menu).", bold_prefix="2. Pengisian Keranjang: ")
create_bullet("Pada halaman keranjang belanja (/cart), pelanggan menginput nama pemesan (2–30 karakter) dan memilih jam serta menit pengambilan (kelipatan 5 menit, minimal 2 menit dari waktu sekarang).", bold_prefix="3. Konfigurasi Penjadwalan: ")
create_bullet("Pelanggan mengklik tombol 'Pay with QRIS'; aplikasi web mengirim HTTP POST payload JSON ke endpoint /api/checkout.", bold_prefix="4. Pengiriman Permintaan Checkout: ")
create_bullet("Worker memvalidasi format nama dan waktu pengambilan, kemudian membuka transaksi atomik Cloud Firestore.", bold_prefix="5. Verifikasi & Transaksi Firestore: ")
create_bullet("Sistem memeriksa sisa stok terkini di database. Jika stok bernilai 0 atau kurang dari permintaan, Worker membatalkan transaksi dan mengembalikan respon HTTP 409 Conflict.", bold_prefix="6. Evaluasi Stok (Decision): ")
create_bullet("Jika stok mencukupi, sistem memotong kuota stok secara atomik (Soft Allocation) dan membuat dokumen pesanan di orders/{orderId} dengan status 'Pending', stockDeducted = true, order_tag = null, dan expiresAt = 60 detik.", bold_prefix="7. Alokasi Stok Sementara: ")
create_bullet("Worker mengirimkan permintaan pembuatan QRIS ke Midtrans Core API (/v2/charge) dengan parameter custom_expiry 1 menit dan acquirer GoPay.", bold_prefix="8. Permintaan QRIS ke Midtrans: ")
create_bullet("Midtrans mengembalikan respons sukses memuat string/URL QRIS dinamis; Worker meneruskannya ke peramban pelanggan.", bold_prefix="9. Pengiriman Data Pembayaran: ")
create_bullet("Antarmuka pelanggan beralih ke rute /pay/:orderId, merender kode QRIS dinamis, mengaktifkan timer hitung mundur 1 menit, dan memasang listener real-time onSnapshot.", bold_prefix="10. Tampilan Pembayaran & Timer: ")
create_bullet("Pelanggan membuka aplikasi m-banking atau e-wallet (GoPay, OVO, Dana, BCA, dll.), memindai QRIS, dan memasukkan PIN transaksi.", bold_prefix="11. Pemindaian & Otorisasi: ")
create_bullet("Dana ditransfer ke rekening perantara Midtrans; proses checkout pada sisi klien selesai sambil menunggu notifikasi server.", bold_prefix="12. Terminasi Pembayaran Klien: ")

create_image("activity_checkout.png", "Gambar 4.1 Activity Diagram Alur Pre-Order & Checkout QRIS Dinamis", width_inches=5.8)

# ------------------------------------------------------------------------------
# 4.2 Activity Diagram Verifikasi Webhook Pembayaran
# ------------------------------------------------------------------------------
create_p("4.2 Activity Diagram Verifikasi Webhook Pembayaran", style='Heading 2', space_before=10, space_after=4)

create_p(
    "Diagram aktivitas pada Gambar 4.2 memodelkan penanganan sinyal pembayaran nirkontak asinkron via HTTP POST "
    "webhook dari Midtrans ke Cloudflare Worker (/api/webhook). Arsitektur ini dirancang dengan prinsip zero-trust "
    "dan idempoten: setiap payload notifikasi diverifikasi keasliannya menggunakan algoritma kriptografi SHA-512 "
    "terhadap MIDTRANS_SERVER_KEY sebelum basis data diperbarui.",
    space_after=6
)

create_p(
    "Tahap paling krusial dalam diagram ini adalah evaluasi ambang batas waktu (Time-To-Live / TTL 1 menit): "
    "jika pelunasan berhasil diverifikasi sebelum 60 detik, sistem meng-increment nomor antrean harian secara atomik "
    "(counters/payment_YYYY-MM-DD), menerbitkan order_tag definitif (contoh: 'Faiz-1'), mengubah status menjadi 'Paid', "
    "dan menyiarkan pembaruan via listener onSnapshot ke dasbor dapur. Sebaliknya, jika pembayaran tiba melampaui batas "
    "1 menit dan stok telah dikembalikan ke menu oleh auto-sweep, transaksi ditandai sebagai 'Expired' dengan flag "
    "latePayment = true. Hal ini memicu 'Late Payment Alert' pada dasbor admin untuk penanganan manual/refund tanpa "
    "mengakibatkan defisit stok fisik dapur.",
    space_after=6
)

# Table 4.4: Specification of Activity Diagram 4.2
t44_headers = ["Komponen Diagram", "Elemen Pemodelan", "Deskripsi Perilaku dalam Sistem"]
t44_data = [
    ["Partisi Swimlane", "Midtrans Payment Gateway", "Mendeteksi pelunasan perbankan, menyusun payload JSON notifikasi, mengirim HTTP POST."],
    ["Partisi Swimlane", "Cloudflare Worker (/api/webhook)", "Menghitung hash SHA-512, memvalidasi signature, mengevaluasi status transaksi dan TTL."],
    ["Partisi Swimlane", "Basis Data Cloud Firestore", "Menyimpan transaksi status Paid/Expired, mencatat timestamp, memicu broadcast onSnapshot."],
    ["Partisi Swimlane", "Smart Queue Dasbor Admin", "Menangkap event real-time, menaikkan kartu ke prioritas atas (hijau) atau menampilkan Late Alert (merah)."],
    ["Partisi Swimlane", "Layar Pembayaran Pelanggan", "Menangkap status Paid, menghentikan timer, menampilkan tiket sukses dengan nomor order_tag."],
    ["Decision Point", "Signature Valid?", "Membandingkan SHA-512 kalkulasi lokal dengan signature_key. Jika tidak cocok, tolak dengan HTTP 403."],
    ["Decision Point", "Status Settlement / Capture?", "Jika transaksi bernilai gagal/batal/kedaluwarsa, lepaskan pesanan dan kembalikan stok."],
    ["Decision Point", "Pelunasan Melebihi Batas 1 Menit?", "Jika <= 1 menit, terbitkan nomor antrean harian. Jika > 1 menit, tandai Expired latePayment."],
]
create_table(t44_headers, t44_data, col_widths=[1.5, 1.8, 3.2])

create_p("Alur eksekusi rinci verifikasi webhook berjalan melalui sekuens berikut:", space_after=4)
create_bullet("Midtrans mendeteksi status settlement/capture dari bank jaringan dan mengirimkan HTTP POST payload ke /api/webhook.", bold_prefix="1. Notifikasi Asinkron: ")
create_bullet("Worker membaca payload dan menghitung nilai hash: SHA512(order_id + status_code + gross_amount + server_key).", bold_prefix="2. Validasi Kriptografi: ")
create_bullet("Jika signature_key tidak cocok, Worker langsung menolak request dengan kode HTTP 403 Forbidden.", bold_prefix="3. Pencegahan Pemalsuan (Decision): ")
create_bullet("Worker memvalidasi parameter transaksi. Jika status transaksi berupa cancel, expire, atau deny, sistem menjalankan fungsi releaseOrder() untuk mengembalikan stok ke katalog menu.", bold_prefix="4. Penanganan Transaksi Gagal: ")
create_bullet("Jika status transaksi adalah settlement/capture, Worker membuka transaksi atomik Firestore untuk memverifikasi nilai pembayaran (gross_amount == total).", bold_prefix="5. Verifikasi Nominal: ")
create_bullet("Worker membandingkan waktu pelunasan terhadap batas waktu pembuatan pesanan (TTL 1 menit).", bold_prefix="6. Evaluasi Batas Waktu TTL (Decision): ")
create_bullet("ALUR TEPAT WAKTU (<= 1 Menit): Worker membaca dokumen counters/payment_YYYY-MM-DD, meng-increment counter harian, membentuk order_tag unik (misal: 'Budi-1'), memperbarui status pesanan menjadi 'Paid', mencatat paidAt, dan merespon HTTP 200 OK ke Midtrans.", bold_prefix="7. Penerbitan Antrean Harian: ")
create_bullet("ALUR TERLAMBAT (> 1 Menit / Late Payment): Jika pesanan telah kedaluwarsa dan stok dikembalikan oleh auto-sweep, pesanan diperbarui menjadi 'Expired' dengan flag latePayment = true, mencatat latePaymentAmount, tanpa mengurangi stok kembali, dan merespon HTTP 200 OK ke Midtrans.", bold_prefix="8. Penanganan Pembayaran Terlambat: ")
create_bullet("Perubahan dokumen pesanan pada koleksi orders memicu penyiaran event WebSocket Firestore onSnapshot secara instan.", bold_prefix="9. Siaran Reaktif Real-Time: ")
create_bullet("Antarmuka pelanggan menangkap status 'Paid', menghentikan timer hitung mundur, dan merender tiket digital yang memuat order_tag dan jadwal pengambilan.", bold_prefix="10. Sinkronisasi Layar Pelanggan: ")
create_bullet("Dasbor Smart Queue admin otomatis menampilkan kartu pesanan baru di posisi teratas antrean aktif (berwarna hijau). Jika pesanan berstatus late payment, muncul banner peringatan merah untuk ditindaklanjuti secara manual oleh pemilik kantin.", bold_prefix="11. Sinkronisasi Dasbor Dapur: ")

create_image("activity_webhook.png", "Gambar 4.2 Activity Diagram Alur Verifikasi Webhook Midtrans & Penomoran Antrean Harian", width_inches=5.8)

# ------------------------------------------------------------------------------
# 4.3 Activity Diagram Penyiapan Dapur & Pengambilan Makanan di Loket
# ------------------------------------------------------------------------------
create_p("4.3 Activity Diagram Penyiapan Dapur & Pengambilan Makanan di Loket", style='Heading 2', space_before=10, space_after=4)

create_p(
    "Diagram aktivitas pada Gambar 4.3 memodelkan operasional fisik di dapur kantin dan loket penyerahan makanan. "
    "Arsitektur antarmuka admin (apps/admin/src/components/Orders.jsx) mengimplementasikan fitur Smart Queue "
    "yang secara otomatis memprioritaskan pesanan berstatus 'Paid' di atas pesanan lainnya dan mengurutkannya "
    "secara kronologis berdasarkan urgensi jam pengambilan terdekat (pickup time slot).",
    space_after=6
)

create_p(
    "Dengan alur kerja ini, koki kantin dapat memasak dan membungkus makanan secara terencana sebelum bel istirahat "
    "berbunyi. Makanan yang siap dibungkus diberi label order_tag (contoh: 'Budi-1') dan diletakkan pada rak "
    "penyerahan loket. Saat jam istirahat tiba, siswa cukup menunjukkan tiket digital atau menyebutkan nomor antrean, "
    "mengambil makanan secara instan tanpa mengantre di kasir (Zero-Wait Pickup), dan petugas menekan tombol 'Complete' "
    "untuk mengakhiri siklus hidup pesanan.",
    space_after=6
)

# Table 4.5: Specification of Activity Diagram 4.3
t45_headers = ["Komponen Diagram", "Elemen Pemodelan", "Deskripsi Perilaku dalam Sistem"]
t45_data = [
    ["Partisi Swimlane", "Petugas / Koki Kantin", "Melihat pesanan pada Smart Queue, memasak makanan, menempelkan label order_tag, menyerahkan paket."],
    ["Partisi Swimlane", "Pelanggan (Siswa / Guru)", "Mendatangi loket saat bel istirahat, menunjukkan tiket digital pada smartphone."],
    ["Partisi Swimlane", "Loket Pengambilan Kantin", "Mencocokkan identitas antrean, serah terima makanan fisik nirkasir, menekan tombol Complete."],
    ["Partisi Swimlane", "Sistem Cloud Firestore & Web", "Memperbarui status menjadi Completed, mencatat completedAt, menyiarkan onSnapshot ke klien."],
    ["Action Node", "Smart Queue Ordering", "Mengelompokkan pesanan Paid berdasarkan batch interval 5 menit terdekat untuk efisiensi produksi dapur."],
    ["Action Node", "Mark Completed Action", "Mengubah state pesanan di database dan memindahkan kartu dari antrean aktif ke riwayat pesanan selesai."],
]
create_table(t45_headers, t45_data, col_widths=[1.5, 1.8, 3.2])

create_p("Rincian urutan aktivitas fisik dan sistem pada alur penyerahan pesanan adalah sebagai berikut:", space_after=4)
create_bullet("Dasbor Smart Queue admin menampilkan kartu pesanan berstatus Paid yang disortir berdasarkan jam istirahat terdekat.", bold_prefix="1. Pemantauan Antrean Dapur: ")
create_bullet("Staf dapur membaca rincian komposisi makanan dan catatan khusus (notes) yang diminta oleh pelanggan.", bold_prefix="2. Pembacaan Rincian Pesanan: ")
create_bullet("Koki memasak dan membungkus makanan sesuai standar porsi dan kebersihan secara terjadwal.", bold_prefix="3. Pengolahan & Pengemasan: ")
create_bullet("Petugas menuliskan atau menempelkan kode antrean harian (order_tag, contoh: 'Budi-1') pada kemasan makanan.", bold_prefix="4. Pelabelan Identitas Antrean: ")
create_bullet("Paket makanan diletakkan pada rak penyerahan cepat di depan loket kantin.", bold_prefix="5. Penempatan di Rak Pengambilan: ")
create_bullet("Bel istirahat sekolah berbunyi; pelanggan mendatangi counter loket Kantin Lina.", bold_prefix="6. Kedatangan Siswa di Counter: ")
create_bullet("Pelanggan menunjukkan tiket digital di smartphone miliknya atau menyebutkan kode antrean ('Budi-1').", bold_prefix="7. Verifikasi Identitas: ")
create_bullet("Petugas mencocokkan kode pada layar siswa dengan label pada kemasan makanan, lalu menyerahkan makanan secara langsung.", bold_prefix="8. Serah Terima Fisik Makanan: ")
create_bullet("Pelanggan menerima makanan dengan waktu tunggu nol (zero waiting time) tanpa perlu melakukan transaksi uang tunai.", bold_prefix="9. Penerimaan Tanpa Antre Kasir: ")
create_bullet("Petugas mengklik tombol aksi 'Complete' berwarna hijau pada kartu pesanan di aplikasi web admin.", bold_prefix="10. Konfirmasi Penyelesaian: ")
create_bullet("Aplikasi web mengeksekusi updateDoc() ke Firestore: status diubah menjadi 'Completed' dan completedAt diisi timestamp terkini.", bold_prefix="11. Pembaruan Status Firestore: ")
create_bullet("Kartu pesanan berpindah secara otomatis dari antrean aktif ke panel riwayat pesanan selesai (Completed Orders History).", bold_prefix="12. Pemindahan ke Riwayat: ")
create_bullet("Listener onSnapshot pada peramban pelanggan mendeteksi status Completed dan mengubah tampilan layar menjadi pesan ramah: 'Order picked up. Enjoy your meal!'.", bold_prefix="13. Konfirmasi Layar Pelanggan: ")

create_image("activity_pickup.png", "Gambar 4.3 Activity Diagram Alur Penyiapan Dapur & Pengambilan Makanan di Loket", width_inches=5.8)

# ------------------------------------------------------------------------------
# 4.4 Activity Diagram Manajemen Menu & Stok Harian
# ------------------------------------------------------------------------------
create_p("4.4 Activity Diagram Manajemen Menu & Stok Harian", style='Heading 2', space_before=10, space_after=4)

create_p(
    "Diagram aktivitas pada Gambar 4.4 memodelkan alur kerja pengelolaan katalog produk makanan dan minuman oleh "
    "Pemilik Kantin Lina melalui antarmuka web admin (apps/admin/src/components/MenuManager.jsx). Seluruh operasi "
    "CRUD (Create, Read, Update, Delete) dilindungi oleh sistem autentikasi Firebase Auth dan terhubung secara "
    "reaktif ke koleksi menu di Cloud Firestore.",
    space_after=6
)

create_p(
    "Fitur ini dirancang khusus dengan memperhatikan kemudahan operasional pemilik kantin: kategori diklasifikasikan "
    "ke dalam 3 pilihan baku ('Food', 'Drink', 'Other'), rasio foto produk distandarisasi persegi (1:1), serta "
    "input harga menerapkan sistem denumerasi ribuan (satuan 'k'). Pemilik kantin cukup mengetik angka '10' untuk "
    "menetapkan harga Rp 10.000, dan aplikasi secara otomatis mengalikannya dengan 1.000 sebelum disimpan ke database. "
    "Setiap perubahan yang disimpan langsung disinkronkan secara live ke peramban pelanggan tanpa perlu refresh.",
    space_after=6
)

# Table 4.6: Specification of Activity Diagram 4.4
t46_headers = ["Komponen Diagram", "Elemen Pemodelan", "Deskripsi Perilaku dalam Sistem"]
t46_data = [
    ["Partisi Swimlane", "Admin / Pemilik Kantin", "Login kredensial, memilih tab Menu & Stock, memilih operasi Tambah/Edit/Hapus, mengisi form."],
    ["Partisi Swimlane", "Aplikasi Web Admin (React 18)", "Mendenumerasi harga (price / 1000), mengalikan input (price * 1000), validasi form, panggil Firestore API."],
    ["Partisi Swimlane", "Basis Data Cloud Firestore", "Mengeksekusi operasi setDoc, updateDoc, atau deleteDoc pada dokumen menu/{id}, broadcast event onSnapshot."],
    ["Partisi Swimlane", "Sinkronisasi Antarmuka", "Memperbarui tampilan katalog admin dan katalog pelanggan secara real-time nirkedip (zero-refresh)."],
    ["Decision Point", "Pilih Jenis Operasi", "Percabangan tiga arah: Tambah Menu Baru, Perbarui Menu / Edit Stok, atau Hapus Menu."],
    ["Decision Point", "Admin Setuju Hapus?", "Dialog konfirmasi peramban (confirm). Jika Batal, proses dihentikan. Jika OK, dokumen dihapus permanen."],
]
create_table(t46_headers, t46_data, col_widths=[1.5, 1.8, 3.2])

create_p("Rincian alur kerja operasional pengelolaan menu mencakup tiga cabang operasi:", space_after=4)
create_bullet("OPERASI TAMBAH MENU: Admin membuka tab Menu & Stock, menginput nama menu, memilih kategori (Food, Drink, Other), mengetik harga dalam ribuan (k, misal ketik '12' untuk Rp 12.000), memasukkan jumlah stok awal, dan menempelkan URL gambar 1:1. Saat tombol 'Tambah Item' diklik, React mengalikan harga dengan 1.000 (Math.round(price * 1000)), membuat ID unik berbasis nama dan timestamp, lalu menuliskan dokumen ke menu/{id} via setDoc().", bold_prefix="1. Tambah Menu Baru: ")
create_bullet("OPERASI EDIT MENU: Admin mengklik tombol 'Edit' pada kartu item yang hendak diubah. Sistem memuat data eksisting ke form dan membagi nominal harga dengan 1.000 (contoh: Rp 10.000 ditampilkan sebagai '10'). Admin memperbarui kuota stok atau mengganti foto, lalu mengklik 'Update Item'. Sistem menyimpan data terbaru ke Firestore via updateDoc().", bold_prefix="2. Perbarui Menu & Stok: ")
create_bullet("OPERASI HAPUS MENU: Admin mengklik tombol 'Delete' berwarna merah pada kartu item. Peramban memunculkan dialog konfirmasi window.confirm('Are you sure you want to delete this item?'). Jika disetujui, fungsi deleteDoc() dieksekusi di Firestore dan dokumen menu terhapus secara permanen.", bold_prefix="3. Penghapusan Menu: ")
create_bullet("SINKRONISASI REAL-TIME: Setiap aksi mutasi data pada koleksi menu memicu event listener onSnapshot pada aplikasi web pelanggan. Menu yang baru ditambah atau stok yang diperbarui langsung tampak secara instan pada layar pelanggan yang sedang aktif berbelanja.", bold_prefix="4. Propagasi Data Live: ")

create_image("activity_menu.png", "Gambar 4.4 Activity Diagram Alur Pengelolaan Katalog Menu & Stok Harian", width_inches=5.8)

# ------------------------------------------------------------------------------
# 4.5 Analisis Sinergi Alur Kerja dan Integritas Sistem
# ------------------------------------------------------------------------------
create_p("4.5 Analisis Sinergi Alur Kerja dan Integritas Sistem", style='Heading 2', space_before=10, space_after=4)

create_p(
    "Keempat Activity Diagram yang telah dimodelkan saling mengunci (interlocking) dan bersinergi dalam membentuk "
    "sebuah sistem pemesanan cloud-native yang tangguh. Melalui pendekatan Reverse Engineering terhadap codebase "
    "Kantin Lina, teridentifikasi sejumlah mekanisme kunci yang menjamin integritas bisnis dan kepuasan pengguna:",
    space_after=6
)

create_bullet(
    "Sistem menerapkan mekanisme Soft Allocation saat checkout untuk mengunci stok secara atomik, namun "
    "membatasi masa tunggu pembayaran tepat 1 menit (60 detik) dengan grace period 0 detik. Mekanisme ini "
    "mengeliminasi risiko perebutan kuota (race condition) antar pelanggan tanpa khawatir terjadi penahanan stok "
    "permanen jika transaksi ditinggalkan.",
    bold_prefix="1. Pencegahan Overbooking & Stok Menggantung: "
)

create_bullet(
    "Nomor urut antrean harian (order_tag) sengaja dipisahkan dari proses checkout dan hanya diterbitkan pada "
    "tahap verifikasi webhook pelunasan. Pendekatan ini menjamin pembukuan kas harian kantin memiliki penomoran yang "
    "runtut tanpa adanya nomor antrean bolong atau fiktif.",
    bold_prefix="2. Eliminasi Nomor Antrean Fiktif: "
)

create_bullet(
    "Jika pelanggan membayar terlambat melampaui jendela 1 menit setelah stok terlanjur dikembalikan oleh sistem "
    "auto-sweep, webhook secara otomatis menandai pesanan sebagai 'Expired' dengan atribut latePayment = true dan "
    "memunculkan banner merah mencolok di dasbor admin. Hal ini melindungi koki dari defisit bahan baku fisik "
    "sekaligus memastikan hak pelanggan tetap tercatat untuk proses refund manual yang akuntabel.",
    bold_prefix="3. Penanganan Anomali Pembayaran Terlambat: "
)

create_bullet(
    "Pelanggan tidak dibebani keharusan registrasi atau login akun pada peramban smartphone (Zero-Friction Ordering), "
    "sedangkan penyerahan makanan fisik di loket berlangsung dalam hitungan detik tanpa antrean kasir (Zero-Wait "
    "Physical Pickup) berkat fitur Smart Queue yang memandu dapur menyiapkan pesanan sebelum bel istirahat tiba.",
    bold_prefix="4. Optimalisasi Pengalaman Pengguna (UX): "
)

print("Bab 4 content insertion complete.")

# 2. Save modified document
doc.save(doc_path)
print(f"Successfully saved updated document to: {doc_path}")
